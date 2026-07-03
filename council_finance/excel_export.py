"""
Excel workbook exports using openpyxl.
Produces a multi-sheet workbook covering all council finance data.
"""
import io
from datetime import datetime, date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              numbers as xl_numbers)
from openpyxl.utils import get_column_letter

# ── Palette ────────────────────────────────────────────────────────────────
NAVY = '1A3A5C'
NAVY_LIGHT = 'D9E0EA'
GREEN_BG = 'D4EDDA'
RED_BG = 'F8D7DA'
AMBER_BG = 'FFF3CD'
WHITE = 'FFFFFF'
GREY = 'F5F7FA'
RECEIPT_BG = 'E6F4EC'
PAYMENT_BG = 'FDF0F0'

MONEY_FMT = '#,##0.00'
DATE_FMT = 'DD/MM/YYYY'

THIN_BORDER = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC'),
)


def _header_fill(hex_colour=NAVY):
    return PatternFill('solid', fgColor=hex_colour)


def _fill(hex_colour):
    return PatternFill('solid', fgColor=hex_colour)


def _col_header(ws, row, col, text, bold=True, colour=NAVY, fg=WHITE):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(bold=bold, color=fg, size=10)
    cell.fill = _header_fill(colour)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = THIN_BORDER
    return cell


def _cell(ws, row, col, value, bold=False, fill_colour=None, number_format=None,
          align='left'):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, size=10)
    if fill_colour:
        cell.fill = _fill(fill_colour)
    if number_format:
        cell.number_format = number_format
    cell.alignment = Alignment(horizontal=align, vertical='center')
    cell.border = THIN_BORDER
    return cell


def _money(ws, row, col, value, bold=False, fill_colour=None):
    try:
        v = float(value)
    except (TypeError, ValueError):
        v = 0.0
    return _cell(ws, row, col, v, bold=bold, fill_colour=fill_colour,
                 number_format=MONEY_FMT, align='right')


def _auto_width(ws, min_w=8, max_w=50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_w, min(max_len + 2, max_w))


def _section_title(ws, row, text, cols, colour=NAVY):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = _header_fill(colour)
    cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    ws.row_dimensions[row].height = 20
    return cell


# ── Cover sheet ─────────────────────────────────────────────────────────────

def _make_cover(wb, fy, council_name, accounting_basis, stats):
    ws = wb.active
    ws.title = 'Summary'
    ws.sheet_view.showGridLines = False

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 25

    # Title block
    title_cell = ws.cell(row=1, column=1, value=council_name)
    title_cell.font = Font(bold=True, size=16, color=NAVY)
    ws.merge_cells('A1:B1')

    sub_cell = ws.cell(row=2, column=1, value=f'Annual Accounts — {fy.label}')
    sub_cell.font = Font(size=12, color=NAVY)
    ws.merge_cells('A2:B2')

    ws.cell(row=3, column=1, value=f'Accounting basis: {accounting_basis}').font = Font(size=10, italic=True)
    ws.cell(row=4, column=1, value=f'Generated: {datetime.now().strftime("%d/%m/%Y %H:%M")}').font = Font(size=9, color='888888')

    ws.row_dimensions[1].height = 30

    # Summary table
    r = 6
    for label, val in [
        ('Total Receipts / Income', stats.get('total_receipts', 0)),
        ('Total Payments / Expenditure', stats.get('total_payments', 0)),
        ('Net Movement', stats.get('net', 0)),
        ('Transactions recorded', stats.get('txn_count', 0)),
        ('Bank balance (latest reconciliation)', stats.get('bank_balance', 0)),
    ]:
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = Font(size=10)
        lc.fill = _fill(GREY)
        lc.border = THIN_BORDER

        vc = ws.cell(row=r, column=2, value=float(val) if isinstance(val, Decimal) else val)
        if isinstance(val, Decimal):
            vc.number_format = MONEY_FMT
        vc.font = Font(bold=True, size=10)
        vc.fill = _fill(GREY)
        vc.border = THIN_BORDER
        vc.alignment = Alignment(horizontal='right')
        r += 1


# ── Cashbook sheet ───────────────────────────────────────────────────────────

def _make_cashbook(wb, transactions, fy):
    ws = wb.create_sheet('Cashbook')
    ws.freeze_panes = 'A3'

    HEADERS = [
        'Date', 'Reference', 'Type', 'Payee / Payer', 'Description',
        'Budget Heading', 'Net (£)', 'VAT (£)', 'Gross (£)',
        'VAT Recl.', 'S137', 'Bank Account', 'Document', 'Notes', 'Void',
    ]
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 20
    ws.cell(row=1, column=1, value=f'Cashbook — {fy.label}').font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells(f'A1:{get_column_letter(len(HEADERS))}1')

    for c, h in enumerate(HEADERS, 1):
        _col_header(ws, 2, c, h)

    r_receipts = Decimal('0')
    r_payments = Decimal('0')

    for row_i, t in enumerate(transactions, 3):
        bg = None if t.is_void else (RECEIPT_BG if t.transaction_type == 'receipt' else PAYMENT_BG)
        _cell(ws, row_i, 1, t.date, fill_colour=bg, number_format=DATE_FMT)
        _cell(ws, row_i, 2, t.reference, fill_colour=bg)
        _cell(ws, row_i, 3, t.type_label, fill_colour=bg)
        _cell(ws, row_i, 4, t.payee_payer, fill_colour=bg)
        _cell(ws, row_i, 5, t.description, fill_colour=bg)
        _cell(ws, row_i, 6, t.heading.name, fill_colour=bg)
        _money(ws, row_i, 7, t.net_amount, fill_colour=bg)
        _money(ws, row_i, 8, t.vat_amount, fill_colour=bg)
        _money(ws, row_i, 9, t.gross_amount, fill_colour=bg)
        _cell(ws, row_i, 10, 'Yes' if t.vat_reclaimable else '', fill_colour=bg)
        _cell(ws, row_i, 11, 'Yes' if t.section_137 else '', fill_colour=bg)
        _cell(ws, row_i, 12, t.bank_account.name if t.bank_account else '', fill_colour=bg)
        _cell(ws, row_i, 13, t.document_filename or '', fill_colour=bg)
        _cell(ws, row_i, 14, t.notes or '', fill_colour=bg)
        _cell(ws, row_i, 15, 'VOID' if t.is_void else '', fill_colour=bg)

        if not t.is_void:
            if t.transaction_type == 'receipt':
                r_receipts += t.gross_amount
            else:
                r_payments += t.gross_amount

    total_row = len(transactions) + 3
    _cell(ws, total_row, 1, 'TOTALS', bold=True, fill_colour=NAVY_LIGHT)
    for c in range(2, 7):
        _cell(ws, total_row, c, '', fill_colour=NAVY_LIGHT)
    _money(ws, total_row, 7, '', fill_colour=NAVY_LIGHT)
    _cell(ws, total_row, 6, 'Receipts →', bold=True, fill_colour=NAVY_LIGHT)
    _money(ws, total_row, 7, r_receipts, bold=True, fill_colour=RECEIPT_BG)
    _cell(ws, total_row, 8, 'Payments →', bold=True, fill_colour=NAVY_LIGHT)
    _money(ws, total_row, 9, r_payments, bold=True, fill_colour=PAYMENT_BG)

    _auto_width(ws)


# ── Budget vs Actuals ────────────────────────────────────────────────────────

def _make_budget(wb, rows, fy):
    ws = wb.create_sheet('Budget vs Actuals')
    ws.freeze_panes = 'A3'

    HEADERS = ['Category', 'Heading', 'Budget (£)', 'Actual (£)', 'Variance (£)', '% Used']
    ws.cell(row=1, column=1, value=f'Budget vs Actuals — {fy.label}').font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells(f'A1:{get_column_letter(len(HEADERS))}1')

    for c, h in enumerate(HEADERS, 1):
        _col_header(ws, 2, c, h)

    for row_i, row in enumerate(rows, 3):
        rag = row['rag']
        bg = GREEN_BG if rag == 'green' else (AMBER_BG if rag == 'amber' else (RED_BG if rag == 'red' else None))
        _cell(ws, row_i, 1, row['heading'].category_label)
        _cell(ws, row_i, 2, row['heading'].name)
        _money(ws, row_i, 3, row['budget'])
        _money(ws, row_i, 4, row['spent'], fill_colour=bg)
        _money(ws, row_i, 5, row['remaining'],
               fill_colour=RED_BG if row['remaining'] < 0 else None)
        pct = float(row['spent'] / row['budget'] * 100) if row['budget'] > 0 else 0
        _cell(ws, row_i, 6, round(pct, 1), fill_colour=bg, number_format='0.0"%"', align='right')

    _auto_width(ws)


# ── R&P Accounts ─────────────────────────────────────────────────────────────

def _make_rp_accounts(wb, fy, receipt_rows, payment_rows,
                       total_receipts, total_payments,
                       opening_balance, closing_balance, council_name):
    ws = wb.create_sheet('R&P Account')
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 18

    r = 1
    ws.cell(row=r, column=1, value=council_name).font = Font(bold=True, size=14, color=NAVY)
    r += 1
    ws.cell(row=r, column=1, value=f'Receipts & Payments Account — {fy.label}').font = Font(size=11, color=NAVY)
    r += 1
    ws.cell(row=r, column=1, value=f'Period: {fy.start_date.strftime("%d/%m/%Y")} to {fy.end_date.strftime("%d/%m/%Y")}').font = Font(size=9, italic=True)
    r += 2

    def _row(label, amount=None, bold=False, indent=False, bg=None):
        nonlocal r
        lbl = ('    ' if indent else '') + label
        lc = ws.cell(row=r, column=1, value=lbl)
        lc.font = Font(bold=bold, size=10)
        if bg:
            lc.fill = _fill(bg)
        if amount is not None:
            vc = ws.cell(row=r, column=2, value=float(amount))
            vc.number_format = MONEY_FMT
            vc.alignment = Alignment(horizontal='right')
            vc.font = Font(bold=bold, size=10)
            if bg:
                vc.fill = _fill(bg)
        r += 1

    _section_title(ws, r, 'RECEIPTS', 2)
    r += 1
    for row in receipt_rows:
        _row(row['heading'].name, row['total'], indent=True)
    _row('Total Receipts', total_receipts, bold=True, bg=RECEIPT_BG)
    r += 1

    _section_title(ws, r, 'PAYMENTS', 2)
    r += 1
    for row in payment_rows:
        _row(row['heading'].name, row['total'], indent=True)
    _row('Total Payments', total_payments, bold=True, bg=PAYMENT_BG)
    r += 1

    _section_title(ws, r, 'SUMMARY', 2, colour='2D5986')
    r += 1
    _row('Opening balance (brought forward)', opening_balance)
    _row('Add: Total Receipts', total_receipts)
    _row('Less: Total Payments', f'({total_payments})')
    _row('Closing balance (carried forward)', closing_balance, bold=True, bg=NAVY_LIGHT)
    r += 1
    ws.cell(row=r, column=1, value='This is an unaudited management account.').font = Font(size=8, italic=True, color='888888')

    _auto_width(ws)


# ── I&E Account ──────────────────────────────────────────────────────────────

def _make_ie_accounts(wb, fy, income_rows, expenditure_rows,
                       total_income, total_expenditure,
                       accruals, council_name):
    ws = wb.create_sheet('I&E Account')
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 18

    r = 1
    ws.cell(row=r, column=1, value=council_name).font = Font(bold=True, size=14, color=NAVY)
    r += 1
    ws.cell(row=r, column=1, value=f'Income & Expenditure Account — {fy.label}').font = Font(size=11, color=NAVY)
    r += 2

    def _row(label, amount=None, bold=False, indent=False, bg=None):
        nonlocal r
        lbl = ('    ' if indent else '') + label
        lc = ws.cell(row=r, column=1, value=lbl)
        lc.font = Font(bold=bold, size=10)
        if bg:
            lc.fill = _fill(bg)
        if amount is not None:
            vc = ws.cell(row=r, column=2, value=float(amount) if isinstance(amount, Decimal) else amount)
            vc.number_format = MONEY_FMT
            vc.alignment = Alignment(horizontal='right')
            vc.font = Font(bold=bold, size=10)
            if bg:
                vc.fill = _fill(bg)
        r += 1

    # Income
    _section_title(ws, r, 'INCOME', 2)
    r += 1
    for row in income_rows:
        _row(row['heading'].name, row['total'], indent=True)

    # Accrued income adjustments
    ai = [a for a in accruals if a.accrual_type in ('accrued_income', 'debtor') and not a.is_reversed]
    if ai:
        _row('Accrued Income (year-end adjustments)', bold=True)
        for a in ai:
            _row(a.description, a.amount, indent=True)
    pi = [a for a in accruals if a.accrual_type == 'prepayment' and not a.is_reversed]
    if pi:
        _row('Less: Prepayments', bold=True)
        for a in pi:
            _row(a.description, -a.amount, indent=True)

    _row('Total Income', total_income, bold=True, bg=RECEIPT_BG)
    r += 1

    # Expenditure
    _section_title(ws, r, 'EXPENDITURE', 2)
    r += 1
    for row in expenditure_rows:
        _row(row['heading'].name, row['total'], indent=True)

    ae = [a for a in accruals if a.accrual_type in ('accrued_expenditure', 'creditor') and not a.is_reversed]
    if ae:
        _row('Accrued Expenditure (year-end adjustments)', bold=True)
        for a in ae:
            _row(a.description, a.amount, indent=True)

    _row('Total Expenditure', total_expenditure, bold=True, bg=PAYMENT_BG)
    r += 1

    surplus = total_income - total_expenditure
    bg = GREEN_BG if surplus >= 0 else RED_BG
    _section_title(ws, r, 'SURPLUS / (DEFICIT)', 2, colour='2D5986')
    r += 1
    _row('Surplus / (Deficit) for the year', surplus, bold=True, bg=bg)

    _auto_width(ws)


# ── Balance Sheet (I&E) ──────────────────────────────────────────────────────

def _make_balance_sheet(wb, fy, assets, accruals, reserves, bank_balance, surplus, council_name):
    ws = wb.create_sheet('Balance Sheet')
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 18

    r = 1
    ws.cell(row=r, column=1, value=council_name).font = Font(bold=True, size=14, color=NAVY)
    r += 1
    ws.cell(row=r, column=1, value=f'Balance Sheet as at {fy.end_date.strftime("%d/%m/%Y")}').font = Font(size=11, color=NAVY)
    r += 2

    def _row(label, amount=None, bold=False, indent=False, bg=None):
        nonlocal r
        lbl = ('    ' if indent else '') + label
        lc = ws.cell(row=r, column=1, value=lbl)
        lc.font = Font(bold=bold, size=10)
        if bg:
            lc.fill = _fill(bg)
        if amount is not None:
            vc = ws.cell(row=r, column=2, value=float(amount) if isinstance(amount, Decimal) else amount)
            vc.number_format = MONEY_FMT
            vc.alignment = Alignment(horizontal='right')
            vc.font = Font(bold=bold, size=10)
            if bg:
                vc.fill = _fill(bg)
        r += 1

    # Fixed assets
    if assets:
        _section_title(ws, r, 'FIXED ASSETS', 2)
        r += 1
        total_cost = Decimal('0')
        total_dep = Decimal('0')
        for a in assets:
            if not a.is_disposed:
                dep = a.accumulated_depreciation(fy.end_date)
                nbv = a.net_book_value(fy.end_date)
                _row(f'{a.name} ({a.category})', nbv, indent=True)
                total_cost += Decimal(str(a.acquisition_cost))
                total_dep += dep
        _row('Total Fixed Assets', total_cost - total_dep, bold=True, bg=NAVY_LIGHT)
        r += 1

    # Current assets
    _section_title(ws, r, 'CURRENT ASSETS', 2)
    r += 1
    _row('Cash and bank', bank_balance, indent=True)
    current_assets = Decimal(str(bank_balance))

    asset_accruals = [a for a in accruals if a.is_asset and not a.is_reversed]
    for a in asset_accruals:
        _row(f'{a.type_label}: {a.description}', a.amount, indent=True)
        current_assets += Decimal(str(a.amount))
    _row('Total Current Assets', current_assets, bold=True)
    r += 1

    # Current liabilities
    liability_accruals = [a for a in accruals if not a.is_asset and not a.is_reversed]
    if liability_accruals:
        _section_title(ws, r, 'CURRENT LIABILITIES', 2)
        r += 1
        total_liabilities = Decimal('0')
        for a in liability_accruals:
            _row(f'{a.type_label}: {a.description}', -a.amount, indent=True)
            total_liabilities += Decimal(str(a.amount))
        _row('Total Current Liabilities', -total_liabilities, bold=True)
        r += 1

    # Net assets
    _section_title(ws, r, 'NET ASSETS', 2, colour='2D5986')
    r += 1
    total_assets = (sum(a.net_book_value(fy.end_date) for a in assets if not a.is_disposed)
                    + current_assets
                    - (total_liabilities if liability_accruals else Decimal('0')))
    _row('Total Net Assets', total_assets, bold=True, bg=NAVY_LIGHT)
    r += 1

    # Represented by
    _section_title(ws, r, 'REPRESENTED BY', 2, colour='2D5986')
    r += 1
    total_reserves = Decimal('0')
    for res in reserves:
        _row(res.name, res.balance, indent=True)
        total_reserves += res.balance
    _row('Surplus / (Deficit) for year', surplus, indent=True)
    _row('Total', total_reserves + surplus, bold=True, bg=NAVY_LIGHT)

    _auto_width(ws)


# ── VAT sheet ────────────────────────────────────────────────────────────────

def _make_vat(wb, vat_txns, fy, receipts_vat, payments_vat, reclaimable, net):
    ws = wb.create_sheet('VAT (126)')
    ws.freeze_panes = 'A3'

    HEADERS = ['Date', 'Reference', 'Type', 'Payee/Payer', 'Description', 'Net (£)', 'VAT (£)', 'Reclaimable']
    ws.cell(row=1, column=1, value=f'VAT Summary — {fy.label}').font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells(f'A1:{get_column_letter(len(HEADERS))}1')

    for c, h in enumerate(HEADERS, 1):
        _col_header(ws, 2, c, h)

    for row_i, t in enumerate(vat_txns, 3):
        _cell(ws, row_i, 1, t.date, number_format=DATE_FMT)
        _cell(ws, row_i, 2, t.reference)
        _cell(ws, row_i, 3, t.type_label)
        _cell(ws, row_i, 4, t.payee_payer)
        _cell(ws, row_i, 5, t.description)
        _money(ws, row_i, 6, t.net_amount)
        _money(ws, row_i, 7, t.vat_amount)
        _cell(ws, row_i, 8, 'Yes' if t.vat_reclaimable else 'No')

    tr = len(vat_txns) + 4
    _cell(ws, tr, 1, 'VAT on receipts:', bold=True)
    _money(ws, tr, 2, receipts_vat, bold=True)
    _cell(ws, tr + 1, 1, 'VAT on payments:', bold=True)
    _money(ws, tr + 1, 2, payments_vat, bold=True)
    _cell(ws, tr + 2, 1, 'Reclaimable VAT:', bold=True)
    _money(ws, tr + 2, 2, reclaimable, bold=True)
    _cell(ws, tr + 3, 1, 'Net reclaimable (VAT126):', bold=True)
    _money(ws, tr + 3, 2, net, bold=True, fill_colour=GREEN_BG)

    _auto_width(ws)


# ── Asset Register ───────────────────────────────────────────────────────────

def _make_assets(wb, assets):
    if not assets:
        return
    ws = wb.create_sheet('Asset Register')
    ws.freeze_panes = 'A3'

    HEADERS = ['Name', 'Category', 'Location', 'Acquired', 'Cost (£)',
               'Depreciation Method', 'Life (yrs)', 'Accumulated Dep (£)',
               'Net Book Value (£)', 'Insurance Value (£)', 'Disposed']
    ws.cell(row=1, column=1, value='Fixed Asset Register').font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells(f'A1:{get_column_letter(len(HEADERS))}1')

    for c, h in enumerate(HEADERS, 1):
        _col_header(ws, 2, c, h)

    for row_i, a in enumerate(assets, 3):
        bg = GREY if a.is_disposed else None
        _cell(ws, row_i, 1, a.name, fill_colour=bg)
        _cell(ws, row_i, 2, a.category, fill_colour=bg)
        _cell(ws, row_i, 3, a.location or '', fill_colour=bg)
        _cell(ws, row_i, 4, a.acquisition_date, fill_colour=bg, number_format=DATE_FMT)
        _money(ws, row_i, 5, a.acquisition_cost, fill_colour=bg)
        _cell(ws, row_i, 6, a.depreciation_method, fill_colour=bg)
        _cell(ws, row_i, 7, a.useful_life_years or '', fill_colour=bg, align='right')
        _money(ws, row_i, 8, a.accumulated_depreciation(), fill_colour=bg)
        _money(ws, row_i, 9, a.net_book_value(), fill_colour=bg)
        _money(ws, row_i, 10, a.insurance_value or 0, fill_colour=bg)
        _cell(ws, row_i, 11, 'Yes' if a.is_disposed else 'No', fill_colour=bg)

    _auto_width(ws)


# ── Reconciliation summary ────────────────────────────────────────────────────

def _make_reconciliations(wb, recons, fy):
    ws = wb.create_sheet('Reconciliations')
    HEADERS = ['Month', 'Bank Account', 'Opening (£)', 'Closing (£)',
               'Outst. Receipts (£)', 'Outst. Payments (£)', 'Difference (£)', 'Status']
    ws.cell(row=1, column=1, value=f'Bank Reconciliations — {fy.label}').font = Font(bold=True, size=12, color=NAVY)
    ws.merge_cells(f'A1:{get_column_letter(len(HEADERS))}1')

    for c, h in enumerate(HEADERS, 1):
        _col_header(ws, 2, c, h)

    for row_i, recon in enumerate(recons, 3):
        bg = GREEN_BG if recon.is_complete else (AMBER_BG if recon.is_balanced else RED_BG)
        _cell(ws, row_i, 1, recon.label)
        _cell(ws, row_i, 2, recon.bank_account.name if recon.bank_account else 'All accounts')
        _money(ws, row_i, 3, recon.opening_balance)
        _money(ws, row_i, 4, recon.closing_balance)
        _money(ws, row_i, 5, recon.outstanding_receipts)
        _money(ws, row_i, 6, recon.outstanding_payments)
        _money(ws, row_i, 7, recon.difference, fill_colour=bg)
        _cell(ws, row_i, 8, 'Signed off' if recon.is_complete else ('Balanced' if recon.is_balanced else 'Not balanced'), fill_colour=bg)

    _auto_width(ws)


# ── Master entry point ────────────────────────────────────────────────────────

def build_workbook(fy, settings, transactions, budget_rows, recons, vat_data,
                   assets, accruals, reserves):
    """
    Build and return an openpyxl Workbook for the given financial year.
    Returns (Workbook, bytes).
    """
    from decimal import Decimal

    wb = Workbook()
    council_name = settings.council_name if settings else 'Council'
    accounting_basis = settings.basis_label if settings else 'Receipts & Payments'

    active_txns = [t for t in transactions if not t.is_void]
    total_receipts = sum(t.gross_amount for t in active_txns if t.transaction_type == 'receipt')
    total_payments = sum(t.gross_amount for t in active_txns if t.transaction_type == 'payment')

    last_recon = next(
        (r for r in sorted(recons, key=lambda x: (x.year, x.month), reverse=True) if r.is_complete),
        None
    )
    bank_balance = Decimal(str(last_recon.closing_balance)) if last_recon else Decimal('0')

    stats = {
        'total_receipts': total_receipts,
        'total_payments': total_payments,
        'net': total_receipts - total_payments,
        'txn_count': len(active_txns),
        'bank_balance': bank_balance,
    }

    _make_cover(wb, fy, council_name, accounting_basis, stats)
    _make_cashbook(wb, sorted(transactions, key=lambda t: (t.date, t.id)), fy)
    _make_budget(wb, budget_rows, fy)

    # R&P or I&E accounts
    receipt_rows = [r for r in budget_rows if r['heading'].is_income]
    payment_rows = [r for r in budget_rows if not r['heading'].is_income]

    first_recon = next(
        (r for r in sorted(recons, key=lambda x: (x.year, x.month)) if r.is_complete),
        None
    )
    opening_balance = Decimal(str(first_recon.opening_balance)) if first_recon else Decimal('0')

    if settings and settings.is_ie:
        ie_income = sum(r['total'] for r in receipt_rows)
        ie_exp = sum(r['total'] for r in payment_rows)
        surplus = ie_income - ie_exp
        _make_ie_accounts(wb, fy, receipt_rows, payment_rows,
                          ie_income, ie_exp, accruals, council_name)
        _make_balance_sheet(wb, fy, assets, accruals, reserves,
                            bank_balance, surplus, council_name)
    else:
        _make_rp_accounts(wb, fy, receipt_rows, payment_rows,
                          total_receipts, total_payments,
                          opening_balance,
                          opening_balance + total_receipts - total_payments,
                          council_name)

    _make_reconciliations(wb, recons, fy)

    vat_txns = vat_data['txns']
    _make_vat(wb, vat_txns, fy,
              vat_data['receipts_vat'], vat_data['payments_vat'],
              vat_data['reclaimable'], vat_data['net'])

    _make_assets(wb, assets)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return wb, output.read()
