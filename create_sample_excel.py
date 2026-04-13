#!/usr/bin/env python3
"""
create_sample_excel.py
Generates a realistic sample orders spreadsheet for testing the invoice generator.

Usage:
    python create_sample_excel.py
    python create_sample_excel.py --output my_orders.xlsx
"""

import argparse
from datetime import date, timedelta
import random

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


SAMPLE_CUSTOMERS = [
    ("Alice Johnson",   "alice.johnson@example.com"),
    ("Bob Smith",       "bob.smith@techcorp.co.uk"),
    ("Clara Williams",  "clara@designstudio.io"),
    ("David Brown",     "david.brown@enterprise.com"),
    ("Eva Martinez",    "eva.martinez@freelance.net"),
]

SAMPLE_PRODUCTS = [
    ("Annual Software Licence – Pro",        299.99),
    ("Professional Consulting (per hour)",    95.00),
    ("Cloud Hosting – Standard Plan",         49.99),
    ("Training Workshop – Full Day",         350.00),
    ("Technical Support Package",            125.00),
    ("Data Analysis Report",                 200.00),
    ("Custom Integration Module",            750.00),
    ("Priority Support (monthly)",            75.00),
    ("Marketing Automation Suite",           180.00),
    ("API Access – Enterprise Tier",         499.00),
]

COLUMNS = [
    "Order Date",
    "Order Number",
    "Customer Name",
    "Email",
    "Item Description",
    "Quantity",
    "Unit Price",
]


def _random_date(start: date, days: int = 90) -> str:
    delta = random.randint(0, days)
    return (start - timedelta(days=delta)).strftime("%d/%m/%Y")


def generate_orders(num_rows: int = 20) -> list[dict]:
    rows = []
    order_counter = 1001
    base = date.today()

    for _ in range(num_rows):
        customer_name, email = random.choice(SAMPLE_CUSTOMERS)
        product, base_price  = random.choice(SAMPLE_PRODUCTS)
        qty   = random.choice([1, 1, 1, 2, 3, 5])
        price = round(base_price * random.uniform(0.95, 1.05), 2)

        rows.append({
            "Order Date":      _random_date(base),
            "Order Number":    f"ORD-{order_counter}",
            "Customer Name":   customer_name,
            "Email":           email,
            "Item Description": product,
            "Quantity":        qty,
            "Unit Price":      price,
        })
        order_counter += 1

    return rows


def write_excel(rows: list[dict], filepath: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"

    # ---- Header style
    header_fill  = PatternFill("solid", fgColor="1A3557")
    header_font  = Font(bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center")
    thin_side    = Side(style="thin", color="CCCCCC")
    thin_border  = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill  = header_fill
        cell.font  = header_font
        cell.alignment = header_align
        cell.border = thin_border

    ws.row_dimensions[1].height = 22

    # ---- Data rows
    alt_fill = PatternFill("solid", fgColor="D6E4F2")
    for row_idx, row in enumerate(rows, start=2):
        fill = alt_fill if row_idx % 2 == 0 else PatternFill()
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[col_name])
            cell.fill   = fill
            cell.border = thin_border
            cell.font   = Font(size=10)

            if col_name in ("Quantity", "Unit Price"):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

    # ---- Column widths
    col_widths = [14, 14, 25, 35, 40, 10, 12]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # ---- Freeze header row
    ws.freeze_panes = "A2"

    wb.save(filepath)
    print(f"Sample orders file written to: {filepath}  ({len(rows)} rows)")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a sample orders Excel file.")
    parser.add_argument(
        "--output", "-o",
        default="sample_orders.xlsx",
        help="Output filename (default: sample_orders.xlsx)",
    )
    parser.add_argument(
        "--rows", "-n",
        type=int,
        default=20,
        help="Number of order rows to generate (default: 20)",
    )
    args = parser.parse_args()

    rows = generate_orders(args.rows)
    write_excel(rows, args.output)


if __name__ == "__main__":
    main()
